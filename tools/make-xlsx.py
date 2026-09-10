# -*- coding: utf-8 -*-
"""
창업진흥원 근무평정 — [별표 9] 점수 보정 실무 계산서 생성 스크립트

개정안 [별표 9](docs/amendments.md §4)의 보정 산식을 엑셀 수식으로 구현한
`보정계산서.xlsx` 를 만든다.

    python3 tools/make-xlsx.py

설계 원칙
  · 값을 하드코딩하지 않는다. 모든 계산은 엑셀 수식으로 넣는다.
  · 함수명은 영문으로 쓴다(한국어 엑셀이 자동 번역한다).
  · 예시 데이터는 모의값이며 창업진흥원의 실제 평정 결과가 아니다.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "보정계산서.xlsx")

# ── 서식 상수 ────────────────────────────────────────────────────────
F_IN   = PatternFill("solid", fgColor="FFF2CC")   # 입력 셀 (연노랑)
F_CALC = PatternFill("solid", fgColor="F2F2F2")   # 계산 셀 (연회색)
F_HEAD = PatternFill("solid", fgColor="D9E1F2")   # 머리글
F_NOTE = PatternFill("solid", fgColor="FCE4D6")   # 고지·경고
B_HEAD = Font(bold=True, size=10)
B_TTL  = Font(bold=True, size=13)
B_SEC  = Font(bold=True, size=11)
FN     = Font(size=10)
FN_S   = Font(size=9, color="595959")
THIN   = Side(style="thin", color="BFBFBF")
BOX    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CTR    = Alignment(horizontal="center", vertical="center")
WRAP   = Alignment(vertical="top", wrap_text=True)

DISCLAIMER = ("고지 — 이 파일의 모든 인명·점수는 산식 검증을 위한 모의 예시 데이터이며, "
              "창업진흥원의 실제 근무평정 결과가 아니다.")

def hdr(ws, row, labels, start=1):
    for k, v in enumerate(labels):
        c = ws.cell(row=row, column=start + k, value=v)
        c.font, c.fill, c.alignment, c.border = B_HEAD, F_HEAD, CTR, BOX
    return row

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def put(ws, addr, val, fill=None, font=FN, align=None, fmt=None):
    c = ws[addr]
    c.value = val
    c.font = font
    if fill: c.fill = fill
    if align: c.alignment = align
    if fmt: c.number_format = fmt
    return c

# ════════════════════════════════════════════════════════════════════
#  예시 데이터 — 평가군 6개 · 직원 37명 (모의)
# ════════════════════════════════════════════════════════════════════
#  평가군 1  1차 · 평가자 1 · 5급 · 10명 (직원01~10)  극단 성향 (60~98)
#  평가군 2  1차 · 평가자 2 · 5급 ·  4명 (직원11~14)  → [별표 9] 제5호
#  평가군 3  1차 · 평가자 3 · 5급 · 16명 (직원15~30)  균등 성향 (78.0~79.5)
#  평가군 4  2차 · 평가자 4 · 5급 · 30명 (직원01~30)
#  평가군 5  1차 · 평가자 5 · 4급 ·  7명 (직원31~37)
#  평가군 6  2차 · 평가자 6 · 4급 ·  7명 (직원31~37)

def emp(i):  return "직원%02d" % i

G1_RAW = [98, 95, 92, 88, 84, 80, 76, 72, 66, 60]                  # 세부지침 §1-2 예제 A
G1_ADJ = [5, 3, 2, 1, 0, 0, -1, -2, -3, -5]                        # 합계 0
G2_RAW = [92, 86, 78, 70]                                          # 세부지침 §1-3 예제 B
G2_ADJ = [4, 1, -1, -4]                                            # 합계 0
G3_RAW = [round(78.0 + 0.1 * k, 1) for k in range(16)]             # 균등 성향 78.0~79.5
G3_ADJ = [0] * 16
# 2차 평가자(본부장)가 5급 30명에게 부여한 원점수 — 1차와 완전히 일치하지 않게 둔다
G4_RAW = [96, 91, 94, 87, 82, 79, 74, 77, 68, 63,
          93, 85, 76, 71, 90, 88, 86, 84, 83, 81,
          80, 78, 75, 73, 72, 70, 69, 67, 65, 61]
G5_RAW = [95, 91, 87, 83, 79, 75, 71]
G5_ADJ = [3, 2, 0, 0, 0, -2, -3]                                   # 합계 0
G6_RAW = [88, 85, 90, 80, 77, 83, 74]                              # 1차와 순위가 다르다

BLOCKS = [
    ("평가군 1", "평가자 1", "1차", "5급", [emp(i) for i in range(1, 11)],  G1_RAW, G1_ADJ),
    ("평가군 2", "평가자 2", "1차", "5급", [emp(i) for i in range(11, 15)], G2_RAW, G2_ADJ),
    ("평가군 3", "평가자 3", "1차", "5급", [emp(i) for i in range(15, 31)], G3_RAW, G3_ADJ),
    ("평가군 4", "평가자 4", "2차", "5급", [emp(i) for i in range(1, 31)],  G4_RAW, None),
    ("평가군 5", "평가자 5", "1차", "4급", [emp(i) for i in range(31, 38)], G5_RAW, G5_ADJ),
    ("평가군 6", "평가자 6", "2차", "4급", [emp(i) for i in range(31, 38)], G6_RAW, None),
]

# 다면평가 — 유효응답 수와 (최고·최저 제외 후) 부여점수 평균. 직원37은 유효응답 2인.
MULTI = {}
for i in range(1, 37):
    MULTI[emp(i)] = (round(70.0 + 0.7 * ((i * 13) % 37), 1), 6 + (i % 5))
MULTI[emp(37)] = ("", 2)
COMMON = {emp(i): 90 + (i * 7) % 11 for i in range(1, 38)}
BOSS = {emp(31): "Y"}          # 부서장(4급) — 상향평가 · 1차:2차 = 50:50 · 현행 배점

wb = Workbook()

SH_USE, SH_GRP, SH_TOT = "사용법", "평가군 계산", "종합점수·등급"
SH_SN, SH_VAL, SH_EX   = "s(n) 표", "검증", "예제"
SH_MIS, SH_LAW         = "MIS 요구사항", "근거 조문"

# ════════════════════════════════════════════════════════════════════
#  ④ s(n) 표  — 다른 시트가 참조하므로 먼저 만든다
# ════════════════════════════════════════════════════════════════════
sn = wb.active
sn.title = SH_SN
SN_R0, SN_N = 5, 39                      # n = 2 ~ 40
SN_R1 = SN_R0 + SN_N - 1                 # 43
COL_S, COL_APPLY, COL_REF, COL_JUDGE = 42, 43, 44, 45   # AP AQ AR AS

put(sn, "A1", "④ s(n) 표 — 인원 n일 때 Blom z값들의 표준편차", font=B_TTL)
put(sn, "A2", "[별표 9] 제3호의 s(n)이다. z가 곧 정규 순서통계량 근사값이므로 s(n)은 그 모표준편차(STDEVP)다. "
              "17명 이상은 [별표 9] 제3호에 따라 s(n) = 1로 본다.", font=FN_S)
put(sn, "A3", "각 셀은 =NORM.S.INV((i−0.375)/(n+0.25)) 이다. 값을 적어 둔 것이 아니라 수식으로 계산한다.", font=FN_S)

hdr(sn, 4, ["n"])
for i in range(1, 41):
    c = sn.cell(row=4, column=1 + i, value=i)
    c.font, c.fill, c.alignment, c.border = B_HEAD, F_HEAD, CTR, BOX
for k, lab in enumerate(["s(n) 계산값 = STDEVP", "규정 적용값 (n≥17 → 1)", "[별표 9] 표 참고값", "일치 판정"]):
    c = sn.cell(row=4, column=COL_S + k, value=lab)
    c.font, c.fill, c.alignment, c.border = B_HEAD, F_HEAD, CTR, BOX

REF_SN = {2: .5895, 3: .7099, 4: .7714, 5: .8097, 6: .8361, 7: .8555, 8: .8704,
          9: .8823, 10: .8920, 11: .9001, 12: .9070, 13: .9129, 14: .9181,
          15: .9226, 16: .9266}
sL, aL = get_column_letter(COL_S), get_column_letter(COL_APPLY)
rL, jL = get_column_letter(COL_REF), get_column_letter(COL_JUDGE)
for k in range(SN_N):
    r = SN_R0 + k
    put(sn, "A%d" % r, k + 2, fill=F_IN, align=CTR)
    for i in range(1, 41):
        cl = get_column_letter(1 + i)
        c = sn.cell(row=r, column=1 + i,
                    value="=IF(%s$4<=$A%d,NORM.S.INV((%s$4-0.375)/($A%d+0.25)),\"\")" % (cl, r, cl, r))
        c.font, c.fill, c.number_format = FN, F_CALC, "0.0000"
    put(sn, "%s%d" % (sL, r), "=STDEVP($B%d:$AO%d)" % (r, r), fill=F_CALC, fmt="0.000000")
    put(sn, "%s%d" % (aL, r), "=IF($A%d>=17,1,%s%d)" % (r, sL, r), fill=F_CALC, fmt="0.0000")
    if k + 2 in REF_SN:
        put(sn, "%s%d" % (rL, r), REF_SN[k + 2], fill=F_CALC, fmt="0.0000")
        put(sn, "%s%d" % (jL, r), '=IF(ROUND(%s%d,4)=%s%d,"PASS","FAIL")' % (sL, r, rL, r),
            fill=F_CALC, align=CTR)
    else:
        put(sn, "%s%d" % (rL, r), "—", fill=F_CALC, align=CTR)
        put(sn, "%s%d" % (jL, r), "참고값 없음 (s=1 적용)", font=FN_S, fill=F_CALC)
sn.conditional_formatting.add("%s%d:%s%d" % (jL, SN_R0, jL, SN_R1),
    CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", fgColor="C6EFCE")))
sn.conditional_formatting.add("%s%d:%s%d" % (jL, SN_R0, jL, SN_R1),
    CellIsRule(operator="equal", formula=['"FAIL"'], fill=PatternFill("solid", fgColor="FFC7CE")))
widths(sn, {"A": 6, sL: 20, aL: 20, rL: 18, jL: 20})
for i in range(1, 41):
    sn.column_dimensions[get_column_letter(1 + i)].width = 9
sn.freeze_panes = "B5"

SN_APPLY = "'%s'!$%s$%d:$%s$%d" % (SH_SN, aL, SN_R0, aL, SN_R1)
SN_CALC  = "'%s'!$%s$%d:$%s$%d" % (SH_SN, sL, SN_R0, sL, SN_R1)
SN_NCOL  = "'%s'!$A$%d:$A$%d" % (SH_SN, SN_R0, SN_R1)

# ════════════════════════════════════════════════════════════════════
#  ② 평가군 계산
# ════════════════════════════════════════════════════════════════════
gp = wb.create_sheet(SH_GRP)
GR0 = 6
rows = []          # (블록 인덱스, 행번호)
r = GR0
block_span = []
for bi, (gname, ev, kind, rank, names, raws, adjs) in enumerate(BLOCKS):
    block_span.append((r, r + len(names) - 1))
    for k in range(len(names)):
        rows.append((bi, r)); r += 1
GR1 = r - 1

put(gp, "A1", "② 평가군 계산 — 개정안 [별표 9] 제1호~제3호·제5호·제7호", font=B_TTL)
put(gp, "A2", DISCLAIMER, fill=F_NOTE, font=FN)
put(gp, "A3", "입력(연노랑): 원점수 · 조정(제7호) · 조정 사유      나머지(연회색)는 모두 자동 계산", font=FN_S)
put(gp, "A4", "동점 입력 건수 →", font=B_SEC)
put(gp, "C4", '=IF(COUNTIF($Q$%d:$Q$%d,"★동점")=0,"0건 (정상)",COUNTIF($Q$%d:$Q$%d,"★동점")&"건 ★ 개정 2 위반 — 순위가 중복되어 산식이 깨진다. 원점수를 고쳐라")'
    % (GR0, GR1, GR0, GR1), fill=F_NOTE, font=Font(bold=True, size=10))
put(gp, "L4", "제5호 적용 평가군 수 →", font=B_SEC)
put(gp, "O4", "=SUMPRODUCT((COUNTIF($A$%d:$A$%d,$A$%d:$A$%d)<6)/COUNTIF($A$%d:$A$%d,$A$%d:$A$%d))"
    % (GR0, GR1, GR0, GR1, GR0, GR1, GR0, GR1), fill=F_CALC, align=CTR)

HEADERS = ["평가군", "평가자명", "평가군 구분", "직급", "피평가자", "원점수",
           "인원 n", "순위 r\n(낮은 순 = 1위)", "s(n)", "p =\n(r−0.375)/(n+0.25)",
           "z = Φ⁻¹(p)", "보정점수\n(제3호)", "적용 기준점수\n(제5호 반영)",
           "조정\n(제7호 ±5)", "조정 후 점수", "조정 사유", "동점 경고",
           "보정 후 순위", "순위 보존 검증", "키 (계산용)",
           "절단 전 보정점수\n(60·100 미적용)", "절단 전 조정 후 점수\n(60·100 미적용)"]
hdr(gp, 5, HEADERS)
gp.row_dimensions[5].height = 34

ALL_A = "$A$%d:$A$%d" % (GR0, GR1)
ALL_F = "$F$%d:$F$%d" % (GR0, GR1)
ALL_L = "$L$%d:$L$%d" % (GR0, GR1)
ALL_T = "$T$%d:$T$%d" % (GR0, GR1)

for (bi, rr) in rows:
    gname, ev, kind, rank, names, raws, adjs = BLOCKS[bi]
    bs, be = block_span[bi]
    k = rr - bs
    blkF = "$F$%d:$F$%d" % (bs, be)
    blkL = "$L$%d:$L$%d" % (bs, be)
    for col, val in (("A", gname), ("B", ev), ("C", kind), ("D", rank), ("E", names[k])):
        put(gp, "%s%d" % (col, rr), val, fill=F_CALC, align=CTR)
    put(gp, "F%d" % rr, raws[k], fill=F_IN, align=CTR, fmt="0.0##")
    put(gp, "G%d" % rr, "=COUNTIF(%s,$A%d)" % (ALL_A, rr), fill=F_CALC, align=CTR)
    put(gp, "H%d" % rr, "=RANK.EQ($F%d,%s,1)" % (rr, blkF), fill=F_CALC, align=CTR)
    put(gp, "I%d" % rr, "=IF($G%d>=17,1,INDEX(%s,MATCH($G%d,%s,0)))" % (rr, SN_APPLY, rr, SN_NCOL),
        fill=F_CALC, align=CTR, fmt="0.0000")
    put(gp, "J%d" % rr, "=($H%d-0.375)/($G%d+0.25)" % (rr, rr), fill=F_CALC, fmt="0.0000")
    put(gp, "K%d" % rr, "=NORM.S.INV($J%d)" % rr, fill=F_CALC, fmt="0.0000")
    put(gp, "L%d" % rr, '=IF($G%d<6,"",ROUND(MEDIAN(60,80+7*$K%d/$I%d,100),3))' % (rr, rr, rr),
        fill=F_CALC, align=CTR, fmt="0.000")
    idx2 = "INDEX(%s,MATCH($E%d&\"|2차\",%s,0))" % (ALL_L, rr, ALL_T)
    put(gp, "M%d" % rr,
        '=IF($G%d>=6,$L%d,IFERROR(IF(%s="",80,%s),80))' % (rr, rr, idx2, idx2),
        fill=F_CALC, align=CTR, fmt="0.000")
    if adjs is None:
        put(gp, "N%d" % rr, 0, fill=F_CALC, align=CTR)
        put(gp, "P%d" % rr, "해당 없음 (제7호 조정은 1차 평가자만)", font=FN_S, fill=F_CALC)
    else:
        put(gp, "N%d" % rr, adjs[k], fill=F_IN, align=CTR)
        put(gp, "P%d" % rr,
            "±5점 조정 — 사유 기재 필수" if abs(adjs[k]) >= 5 else ("조정 사유 기재" if adjs[k] else ""),
            fill=F_IN, align=WRAP)
    put(gp, "O%d" % rr, '=IF($N%d="",$M%d,ROUND(MEDIAN(60,$M%d+$N%d,100),3))' % (rr, rr, rr, rr),
        fill=F_CALC, align=CTR, fmt="0.000")
    put(gp, "Q%d" % rr, '=IF(COUNTIFS(%s,$A%d,%s,$F%d)>1,"★동점","")' % (ALL_A, rr, ALL_F, rr),
        fill=F_CALC, align=CTR)
    put(gp, "R%d" % rr, '=IF($L%d="","",RANK.EQ($L%d,%s,1))' % (rr, rr, blkL), fill=F_CALC, align=CTR)
    put(gp, "S%d" % rr, '=IF($L%d="","제5호 - 미실시",IF($R%d=$H%d,"보존","★불일치"))' % (rr, rr, rr),
        fill=F_CALC, align=CTR)
    put(gp, "T%d" % rr, '=$E%d&"|"&$C%d' % (rr, rr), font=FN_S, fill=F_CALC)
    # 절단이 실제로 발생하는지 세려면 절단 전 값이 있어야 한다(⑤ 검증에서 COUNTIF 로 센다)
    put(gp, "U%d" % rr, '=IF($G%d<6,"",80+7*$K%d/$I%d)' % (rr, rr, rr), fill=F_CALC, align=CTR, fmt="0.000")
    put(gp, "V%d" % rr, '=IF($N%d="",$M%d,$M%d+$N%d)' % (rr, rr, rr, rr), fill=F_CALC, align=CTR, fmt="0.000")

for a in ("Q", "S"):
    gp.conditional_formatting.add("%s%d:%s%d" % (a, GR0, a, GR1),
        CellIsRule(operator="containsText", formula=['NOT(ISERROR(SEARCH("★",%s%d)))' % (a, GR0)],
                   fill=PatternFill("solid", fgColor="FFC7CE")))
widths(gp, {"A": 10, "B": 10, "C": 11, "D": 7, "E": 10, "F": 9, "G": 8, "H": 11, "I": 9,
            "J": 12, "K": 10, "L": 11, "M": 14, "N": 10, "O": 12, "P": 28, "Q": 10,
            "R": 11, "S": 14, "T": 14, "U": 16, "V": 18})
gp.freeze_panes = "F6"

GP  = lambda a: "'%s'!%s" % (SH_GRP, a)
G_G = GP(ALL_A.replace("$A", "$G"))
G_O = GP("$O$%d:$O$%d" % (GR0, GR1))
G_L = GP(ALL_L)
G_T = GP(ALL_T)
G_H = GP("$H$%d:$H$%d" % (GR0, GR1))

# ════════════════════════════════════════════════════════════════════
#  ③ 종합점수·등급
# ════════════════════════════════════════════════════════════════════
tt = wb.create_sheet(SH_TOT)
PR0 = 9
PEOPLE = [emp(i) for i in range(1, 38)]
RANK_OF = {emp(i): ("5급" if i <= 30 else "4급") for i in range(1, 38)}
PR1 = PR0 + len(PEOPLE) - 1        # 45
QR0 = 50                            # 정원 블록 (5급 / 4급)

put(tt, "A1", "③ 종합점수·등급 — 규칙 제9조② · 제26조① · [별표 2] · [별표 7]", font=B_TTL)
put(tt, "A2", DISCLAIMER, fill=F_NOTE, font=FN)
put(tt, "A3", "개정 14 채택 여부 (Y/N) →", font=B_SEC)
put(tt, "B3", "Y", fill=F_IN, align=CTR, font=Font(bold=True, size=11))
put(tt, "C3", "★ 배점 스위치 — Y: 업적·역량 75% / 다면 20% / 공통 5%   N(현행): 85% / 10% / 5%", font=FN_S)
put(tt, "A4", "업적·역량 배점", font=FN)
put(tt, "B4", '=IF($B$3="Y",0.75,0.85)', fill=F_CALC, align=CTR, fmt="0.00")
put(tt, "A5", "다면(상향) 배점", font=FN)
put(tt, "B5", '=IF($B$3="Y",0.2,0.1)', fill=F_CALC, align=CTR, fmt="0.00")
put(tt, "A6", "공통평가 배점", font=FN)
put(tt, "B6", 0.05, fill=F_CALC, align=CTR, fmt="0.00")
put(tt, "C6", "부서장(4급 이상)은 다면 대신 상향평가이며 1차:2차 = 50:50, 배점은 현행(0.85/0.10/0.05)을 유지한다.", font=FN_S)
put(tt, "C7", "[별표 9] 제5호 적용 행(1차 평가군 6명 미만)은 1차 비율 0 · 2차 비율 100%이며, "
              "업적·역량 = 2차 보정점수 × 100% + 1차 평가자의 제7호 조정량(E열 − F열)이다.", font=FN_S)

TH = ["직원", "직급", "부서장\n여부", "1차 평가군\n인원 n₁", "1차 적용점수\n(조정 후)\n제5호 시 = 2차+조정",
      "2차 보정점수", "제5호", "1차 비율", "2차 비율", "업적·역량 점수",
      "다면/상향\n원점수", "유효\n응답수", "다면 평가군\n인원", "다면 순위", "다면 s(n)",
      "다면 보정점수", "공통평가", "업적·역량\n배점", "다면 배점", "종합점수",
      "직급 인원", "직급내 순위", "S 정원", "A 정원", "B 정원", "C 정원",
      "등급", "등급 인원", "등급내 순위", "평정점 (별표 7)"]
hdr(tt, 8, TH)
tt.row_dimensions[8].height = 34

P_K = "$K$%d:$K$%d" % (PR0, PR1)
P_L = "$L$%d:$L$%d" % (PR0, PR1)
P_B = "$B$%d:$B$%d" % (PR0, PR1)
P_T = "$T$%d:$T$%d" % (PR0, PR1)
QA  = "$A$%d:$A$%d" % (QR0, QR0 + 1)

for k, nm in enumerate(PEOPLE):
    rr = PR0 + k
    put(tt, "A%d" % rr, nm, fill=F_CALC, align=CTR)
    put(tt, "B%d" % rr, RANK_OF[nm], fill=F_CALC, align=CTR)
    put(tt, "C%d" % rr, BOSS.get(nm, "N"), fill=F_IN, align=CTR)
    m1 = 'MATCH($A%d&"|1차",%s,0)' % (rr, G_T)
    m2 = 'MATCH($A%d&"|2차",%s,0)' % (rr, G_T)
    put(tt, "D%d" % rr, "=INDEX(%s,%s)" % (G_G, m1), fill=F_CALC, align=CTR)
    put(tt, "E%d" % rr, "=INDEX(%s,%s)" % (G_O, m1), fill=F_CALC, align=CTR, fmt="0.000")
    i2 = "INDEX(%s,%s)" % (G_L, m2)
    put(tt, "F%d" % rr, '=IFERROR(IF(%s="",80,%s),80)' % (i2, i2), fill=F_CALC, align=CTR, fmt="0.000")
    put(tt, "G%d" % rr, '=IF($D%d<6,"제5호 — 2차 100%% + 1차 조정","-")' % rr, fill=F_CALC, align=CTR)
    put(tt, "H%d" % rr, '=IF($D%d<6,0,IF($C%d="Y",0.5,0.6))' % (rr, rr), fill=F_CALC, align=CTR, fmt="0.00")
    put(tt, "I%d" % rr, '=IF($D%d<6,1,IF($C%d="Y",0.5,0.4))' % (rr, rr), fill=F_CALC, align=CTR, fmt="0.00")
    # 제5호 적용 시: 2차 보정점수 100%($I×$F) + 1차 평가자의 제7호 조정량($E−$F)
    put(tt, "J%d" % rr, "=ROUND(IF($D%d<6,$I%d*$F%d+($E%d-$F%d),$H%d*$E%d+$I%d*$F%d),3)"
        % (rr, rr, rr, rr, rr, rr, rr, rr, rr), fill=F_CALC, align=CTR, fmt="0.000")
    mv, mc = MULTI[nm]
    put(tt, "K%d" % rr, mv, fill=F_IN, align=CTR, fmt="0.0")
    put(tt, "L%d" % rr, mc, fill=F_IN, align=CTR)
    put(tt, "M%d" % rr, '=COUNTIFS(%s,">=3")' % P_L, fill=F_CALC, align=CTR)
    put(tt, "N%d" % rr, '=IF($L%d<3,"",COUNTIFS(%s,">=3",%s,"<"&$K%d)+1)' % (rr, P_L, P_K, rr),
        fill=F_CALC, align=CTR)
    put(tt, "O%d" % rr, '=IF($L%d<3,"",IF($M%d>=17,1,INDEX(%s,MATCH($M%d,%s,0))))'
        % (rr, rr, SN_APPLY, rr, SN_NCOL), fill=F_CALC, align=CTR, fmt="0.0000")
    put(tt, "P%d" % rr,
        '=IF($L%d<3,80,ROUND(MEDIAN(60,80+7*NORM.S.INV(($N%d-0.375)/($M%d+0.25))/$O%d,100),3))'
        % (rr, rr, rr, rr), fill=F_CALC, align=CTR, fmt="0.000")
    put(tt, "Q%d" % rr, COMMON[nm], fill=F_IN, align=CTR)
    put(tt, "R%d" % rr, '=IF($C%d="Y",0.85,$B$4)+IF($L%d<3,IF($C%d="Y",0.1,$B$5),0)' % (rr, rr, rr),
        fill=F_CALC, align=CTR, fmt="0.00")
    put(tt, "S%d" % rr, '=IF($L%d<3,0,IF($C%d="Y",0.1,$B$5))' % (rr, rr), fill=F_CALC, align=CTR, fmt="0.00")
    put(tt, "T%d" % rr, "=ROUND($R%d*$J%d+$S%d*$P%d+$B$6*$Q%d,3)" % (rr, rr, rr, rr, rr),
        fill=F_CALC, align=CTR, fmt="0.000")
    put(tt, "U%d" % rr, "=COUNTIF(%s,$B%d)" % (P_B, rr), fill=F_CALC, align=CTR)
    put(tt, "V%d" % rr, '=COUNTIFS(%s,$B%d,%s,">"&$T%d)+1' % (P_B, rr, P_T, rr), fill=F_CALC, align=CTR)
    for ci, qc in zip(("W", "X", "Y", "Z"), ("T", "U", "V", "W")):
        put(tt, "%s%d" % (ci, rr),
            "=INDEX($%s$%d:$%s$%d,MATCH($B%d,%s,0))" % (qc, QR0, qc, QR0 + 1, rr, QA),
            fill=F_CALC, align=CTR)
    put(tt, "AA%d" % rr,
        '=IF($V%d<=$W%d,"S",IF($V%d<=$W%d+$X%d,"A",IF($V%d<=$W%d+$X%d+$Y%d,"B","C")))'
        % (rr, rr, rr, rr, rr, rr, rr, rr, rr), fill=F_CALC, align=CTR, font=Font(bold=True, size=10))
    put(tt, "AB%d" % rr,
        '=IF($AA%d="S",$W%d,IF($AA%d="A",$X%d,IF($AA%d="B",$Y%d,$Z%d)))'
        % (rr, rr, rr, rr, rr, rr, rr), fill=F_CALC, align=CTR)
    put(tt, "AC%d" % rr,
        '=$V%d-IF($AA%d="S",0,IF($AA%d="A",$W%d,IF($AA%d="B",$W%d+$X%d,$W%d+$X%d+$Y%d)))'
        % (rr, rr, rr, rr, rr, rr, rr, rr, rr, rr), fill=F_CALC, align=CTR)
    put(tt, "AD%d" % rr,
        '=ROUND(IF($AA%d="S",100,IF($AA%d="A",90,IF($AA%d="B",80,70)))-(10/$AB%d)*($AC%d-1),3)'
        % (rr, rr, rr, rr, rr), fill=F_CALC, align=CTR, fmt="0.00")

# ── 정원 블록 (최대잔여법) ───────────────────────────────────────────
put(tt, "A48", "직급별 등급 정원 — 최대잔여법 (규칙 제26조①③ · 〈09 수치의 근거〉 §2-3)", font=B_SEC)
QH = ["직급", "인원 m", "S 몫\nm×0.20", "A 몫\nm×0.30", "B 몫\nm×0.40", "C 몫\nm×0.10",
      "S 정수부", "A 정수부", "B 정수부", "C 정수부", "남은 자리",
      "S 소수부", "A 소수부", "B 소수부", "C 소수부",
      "S 소수부순위", "A 소수부순위", "B 소수부순위", "C 소수부순위",
      "S 정원", "A 정원", "B 정원", "C 정원", "정원 합계 검증"]
hdr(tt, 49, QH)
tt.row_dimensions[49].height = 30
for k, rk in enumerate(("5급", "4급")):
    rr = QR0 + k
    put(tt, "A%d" % rr, rk, fill=F_CALC, align=CTR)
    put(tt, "B%d" % rr, "=COUNTIF(%s,$A%d)" % (P_B, rr), fill=F_CALC, align=CTR)
    for ci, ratio in zip(("C", "D", "E", "F"), (0.2, 0.3, 0.4, 0.1)):
        put(tt, "%s%d" % (ci, rr), "=$B%d*%s" % (rr, ratio), fill=F_CALC, align=CTR, fmt="0.00")
    for ci, src in zip(("G", "H", "I", "J"), ("C", "D", "E", "F")):
        put(tt, "%s%d" % (ci, rr), "=INT($%s%d)" % (src, rr), fill=F_CALC, align=CTR)
    put(tt, "K%d" % rr, "=$B%d-SUM($G%d:$J%d)" % (rr, rr, rr), fill=F_CALC, align=CTR)
    for ci, src, iq in zip(("L", "M", "N", "O"), ("C", "D", "E", "F"), ("G", "H", "I", "J")):
        put(tt, "%s%d" % (ci, rr), "=$%s%d-$%s%d" % (src, rr, iq, rr), fill=F_CALC, align=CTR, fmt="0.00")
    for ci, src in zip(("P", "Q", "R", "S"), ("L", "M", "N", "O")):
        put(tt, "%s%d" % (ci, rr), "=RANK.EQ($%s%d,$L%d:$O%d,0)" % (src, rr, rr, rr),
            fill=F_CALC, align=CTR)
    for ci, iq, rk_ in zip(("T", "U", "V", "W"), ("G", "H", "I", "J"), ("P", "Q", "R", "S")):
        put(tt, "%s%d" % (ci, rr), "=$%s%d+IF($%s%d<=$K%d,1,0)" % (iq, rr, rk_, rr, rr),
            fill=F_CALC, align=CTR, font=Font(bold=True, size=10))
    put(tt, "X%d" % rr, '=IF(SUM($T%d:$W%d)=$B%d,"PASS","FAIL")' % (rr, rr, rr),
        fill=F_CALC, align=CTR)
tt.conditional_formatting.add("X%d:X%d" % (QR0, QR0 + 1),
    CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", fgColor="C6EFCE")))
tt.conditional_formatting.add("X%d:X%d" % (QR0, QR0 + 1),
    CellIsRule(operator="equal", formula=['"FAIL"'], fill=PatternFill("solid", fgColor="FFC7CE")))
put(tt, "A53", "① 인원×비율의 정수부를 먼저 배정한다  ② 남은 자리는 소수부가 큰 등급 순으로 하나씩 배정한다. "
               "4급 7명의 결과 S1·A2·B3·C1은 개정안 [별표 10]의 7명 행과 일치한다.", font=FN_S)
widths(tt, {"A": 10, "B": 8, "C": 9, "D": 12, "E": 13, "F": 12, "G": 11, "H": 9, "I": 9,
            "J": 13, "K": 11, "L": 9, "M": 12, "N": 10, "O": 10, "P": 13, "Q": 9,
            "R": 11, "S": 10, "T": 11, "U": 10, "V": 12, "W": 9, "X": 9, "Y": 9, "Z": 9,
            "AA": 7, "AB": 10, "AC": 11, "AD": 14})
tt.freeze_panes = "C9"

# ════════════════════════════════════════════════════════════════════
#  ⑤ 검증
# ════════════════════════════════════════════════════════════════════
vv = wb.create_sheet(SH_VAL)
put(vv, "A1", "⑤ 검증 — 개정안 [별표 9]의 산식이 제4호가 정한 값을 실제로 달성하는가", font=B_TTL)
put(vv, "A2", DISCLAIMER, fill=F_NOTE, font=FN)
put(vv, "A3", "모든 판정은 수식이다. ② 평가군 계산의 원점수를 바꾸면 이 시트가 즉시 다시 판정한다.", font=FN_S)

def sec(row, text):
    put(vv, "A%d" % row, text, font=B_SEC)

sec(5, "1. 평가군별 보정점수 — 평균 80.00 · 표준편차 7.00 (제3호·제4호) / 제7호 조정 합계 0.00")
hdr(vv, 6, ["평가군", "구분", "인원 n", "보정점수 평균", "판정\n(=80.00)",
            "보정점수 표준편차\nSTDEVP", "제3호가 정하는\n기대 표준편차",
            "판정", "÷s(n) 없을 때\n표준편차 = 7×s(n)", "제7호 조정 합계", "판정\n(=0.00)"])
vv.row_dimensions[6].height = 34
vr = 7
for bi, (gname, ev, kind, rank, names, raws, adjs) in enumerate(BLOCKS):
    bs, be = block_span[bi]
    bL = GP("$L$%d:$L$%d" % (bs, be))
    bK = GP("$K$%d:$K$%d" % (bs, be))
    bN = GP("$N$%d:$N$%d" % (bs, be))
    put(vv, "A%d" % vr, gname, fill=F_CALC, align=CTR)
    put(vv, "B%d" % vr, kind, fill=F_CALC, align=CTR)
    put(vv, "C%d" % vr, len(names), fill=F_CALC, align=CTR)
    if len(names) >= 6:
        put(vv, "D%d" % vr, "=ROUND(AVERAGE(%s),2)" % bL, fill=F_CALC, align=CTR, fmt="0.00")
        put(vv, "E%d" % vr, '=IF($D%d=80,"PASS","FAIL")' % vr, fill=F_CALC, align=CTR)
        put(vv, "F%d" % vr, "=ROUND(STDEVP(%s),2)" % bL, fill=F_CALC, align=CTR, fmt="0.00")
        # 제3호 후단(17명 이상은 s(n)=1)을 그대로 적용하면 n>=17 평가군의 기대 표준편차는 7×s(n) 이다.
        put(vv, "G%d" % vr,
            '=IF($C%d<=16,7,IFERROR(ROUND(7*INDEX(%s,MATCH($C%d,%s,0)),2),"④ 표(n≤40) 확장 필요"))'
            % (vr, SN_CALC, vr, SN_NCOL), fill=F_CALC, align=CTR, fmt="0.00")
        put(vv, "H%d" % vr, '=IF($F%d=$G%d,"PASS","FAIL")' % (vr, vr), fill=F_CALC, align=CTR)
        put(vv, "I%d" % vr, "=ROUND(7*STDEVP(%s),2)" % bK, fill=F_CALC, align=CTR, fmt="0.00")
    else:
        for c in ("D", "F", "G", "I"):
            put(vv, "%s%d" % (c, vr), "제5호 미실시", font=FN_S, fill=F_CALC, align=CTR)
        for c in ("E", "H"):
            put(vv, "%s%d" % (c, vr), "해당 없음", font=FN_S, fill=F_CALC, align=CTR)
    put(vv, "J%d" % vr, "=ROUND(SUM(%s),2)" % bN, fill=F_CALC, align=CTR, fmt="0.00")
    put(vv, "K%d" % vr, '=IF($J%d=0,"PASS","FAIL")' % vr, fill=F_CALC, align=CTR)
    vr += 1
V1_R0, V1_R1 = 7, vr - 1
put(vv, "A%d" % vr, "★ I열이 제4호의 목표 7.00에 미달한다는 것이 ÷s(n) 을 두는 이유다. "
                    "F열(실제)과 G열(제3호가 정하는 값)이 일치하면 구현이 규정대로 된 것이다.", font=FN_S)
vr += 2

# ── 논점 — 제3호 후단(17명 이상 s(n)=1)의 효과 ───────────────────────
sec(vr, "1-2. 조합이 먼저 밝히는 논점 — 제3호 후단 「17명 이상은 s(n) = 1로 본다」의 효과")
put(vv, "A%d" % (vr + 1),
    "1차 평가군(부서 단위)은 창업진흥원에서 3~16명이므로 s(n) 표가 적용되어 표준편차가 정확히 7.00이 된다. "
    "그러나 2차 평가군(본부 단위)은 통상 17명 이상이어서 제3호 후단에 따라 s(n)=1 이 적용되고, "
    "표준편차가 7.00에 미달한다. 아래 값은 모두 ④ 시트에서 계산된 것이다.", font=FN, align=WRAP)
vv.merge_cells(start_row=vr + 1, start_column=1, end_row=vr + 1, end_column=8)
vv.row_dimensions[vr + 1].height = 44
hdr(vv, vr + 2, ["평가군 인원 n", "s(n) 실제값 (④ 시트)", "제3호 후단 적용 s(n)",
                 "그때의 표준편차 = 7×s(n)", "목표 7.00과의 차", "s(n) 표를 그대로 쓰면"])
vv.row_dimensions[vr + 2].height = 34
rX = vr + 3
for n in (16, 17, 20, 30, 40):
    put(vv, "A%d" % rX, n, fill=F_CALC, align=CTR)
    put(vv, "B%d" % rX, "=INDEX(%s,MATCH($A%d,%s,0))" % (SN_CALC, rX, SN_NCOL), fill=F_CALC, align=CTR, fmt="0.0000")
    put(vv, "C%d" % rX, "=INDEX(%s,MATCH($A%d,%s,0))" % (SN_APPLY, rX, SN_NCOL), fill=F_CALC, align=CTR, fmt="0.0000")
    put(vv, "D%d" % rX, "=ROUND(7*$B%d/$C%d,2)" % (rX, rX), fill=F_CALC, align=CTR, fmt="0.00")
    put(vv, "E%d" % rX, "=ROUND($D%d-7,2)" % rX, fill=F_CALC, align=CTR, fmt="+0.00;-0.00;0.00")
    put(vv, "F%d" % rX, "=ROUND(7*$B%d/$B%d,2)" % (rX, rX), fill=F_CALC, align=CTR, fmt="0.00")
    rX += 1
put(vv, "A%d" % rX,
    "해소 방법 — ④ 시트는 이미 n = 40까지 s(n)을 수식으로 계산해 두었다. [별표 9] 제3호 후단을 삭제하고 "
    "s(n) 표를 인원 구간 전체로 확대하면 F열처럼 전 구간에서 7.00이 된다(조문 한 줄). "
    "그대로 두는 경우에도 1차 평가군은 영향을 받지 않으며, 2차 평가군의 표준편차가 약 6.5~6.8점이 되어 "
    "2차 평가자의 점수 폭이 1차보다 조금 좁아진다는 뜻이다. 어느 쪽을 택할지는 노사 협의 사항이다.", font=FN_S, align=WRAP)
vv.merge_cells(start_row=rX, start_column=1, end_row=rX, end_column=8)
vv.row_dimensions[rX].height = 58
vr = rX + 2

sec(vr, "2. ÷s(n) 의 효과 — 나눗셈이 없으면 어느 평가군도 제4호의 표준편차 7점을 달성하지 못한다")
hdr(vv, vr + 1, ["평가군 인원 n", "s(n)  (④ 시트 참조)", "÷s(n) 없을 때 표준편차\n= 7 × s(n)",
                 "÷s(n) 적용 시 표준편차\n= 7 × s(n) ÷ s(n)", "개정안 명시값", "판정"])
vv.row_dimensions[vr + 1].height = 34
ABL = {3: 4.97, 5: 5.67, 10: 6.24, 16: 6.49}
r2 = vr + 2
for n, ref in ABL.items():
    put(vv, "A%d" % r2, n, fill=F_CALC, align=CTR)
    put(vv, "B%d" % r2, "=INDEX(%s,MATCH($A%d,%s,0))" % (SN_CALC, r2, SN_NCOL), fill=F_CALC, align=CTR, fmt="0.0000")
    put(vv, "C%d" % r2, "=ROUND(7*$B%d,2)" % r2, fill=F_CALC, align=CTR, fmt="0.00")
    put(vv, "D%d" % r2, "=ROUND(7*$B%d/$B%d,2)" % (r2, r2), fill=F_CALC, align=CTR, fmt="0.00")
    put(vv, "E%d" % r2, ref, fill=F_CALC, align=CTR, fmt="0.00")
    put(vv, "F%d" % r2, '=IF(AND($C%d=$E%d,$D%d=7),"PASS","FAIL")' % (r2, r2, r2), fill=F_CALC, align=CTR)
    r2 += 1
V2_R0, V2_R1 = r2 - len(ABL), r2 - 1
put(vv, "A%d" % r2, "E열은 개정안 [별표 9] ÷s(n) 설명표의 값(4.97 / 5.67 / 6.24 / 6.49)이다. C열이 그 값을 재현한다.", font=FN_S)
r2 += 2

sec(r2, "3. 부서 내 순위 보존 · 절단 · 동점 · s(n) 표")
hdr(vv, r2 + 1, ["검증 항목", "산출값", "요구값", "판정"])
r3 = r2 + 2
G_S = GP("$S$%d:$S$%d" % (GR0, GR1))
G_Q = GP("$Q$%d:$Q$%d" % (GR0, GR1))
G_K = GP("$K$%d:$K$%d" % (GR0, GR1))
G_I = GP("$I$%d:$I$%d" % (GR0, GR1))
G_M = GP("$M$%d:$M$%d" % (GR0, GR1))
G_N = GP("$N$%d:$N$%d" % (GR0, GR1))
G_U = GP("$U$%d:$U$%d" % (GR0, GR1))
G_V = GP("$V$%d:$V$%d" % (GR0, GR1))
SN_J = "'%s'!$%s$%d:$%s$%d" % (SH_SN, jL, SN_R0, jL, SN_R0 + 14)   # n=2~16 만 참고값이 있다
CHECKS = [
    ("원점수 순위 ≠ 보정 후 순위 인 건수 (순위 100% 보존)",
     '=COUNTIF(%s,"★불일치")' % G_S, 0),
    ("동점 입력 건수 (개정 2 — 동점 금지)",
     '=COUNTIF(%s,"★동점")' % G_Q, 0),
    ("보정점수 절단(60점 미만) 발생 건수 — ② U열 기준",
     '=COUNTIF(%s,"<60")' % G_U, 0),
    ("보정점수 절단(100점 초과) 발생 건수 — ② U열 기준",
     '=COUNTIF(%s,">100")' % G_U, 0),
    ("제7호 조정 후 절단 발생 건수 — ② V열 기준",
     '=COUNTIF(%s,"<60")+COUNTIF(%s,">100")' % (G_V, G_V), 0),
    ("s(n) 표(n=2~16) 참고값 불일치 건수",
     '=COUNTIF(%s,"FAIL")' % SN_J, 0),
    ("제7호 조정 총합계 (전 평가군)",
     "=ROUND(SUM(%s),2)" % G_N, 0),
]
for lab, f, want in CHECKS:
    put(vv, "A%d" % r3, lab, font=FN, fill=F_CALC, align=WRAP)
    put(vv, "B%d" % r3, f, fill=F_CALC, align=CTR)
    put(vv, "C%d" % r3, want, fill=F_CALC, align=CTR)
    put(vv, "D%d" % r3, '=IF($B%d=$C%d,"PASS","FAIL")' % (r3, r3), fill=F_CALC, align=CTR)
    r3 += 1
V3_R0, V3_R1 = r3 - len(CHECKS), r3 - 1
r3 += 1

sec(r3, "4. 등급 정원 합계 = 직급 인원 (규칙 제26조① · 최대잔여법) / [별표 10] 대조")
hdr(vv, r3 + 1, ["직급", "직급 인원", "S", "A", "B", "C", "정원 합계", "판정", "[별표 10] 대조"])
r4 = r3 + 2
TOT = lambda a: "'%s'!%s" % (SH_TOT, a)
for k, rk in enumerate(("5급", "4급")):
    qr = QR0 + k
    put(vv, "A%d" % r4, rk, fill=F_CALC, align=CTR)
    put(vv, "B%d" % r4, "=%s" % TOT("$B$%d" % qr), fill=F_CALC, align=CTR)
    for ci, qc in zip(("C", "D", "E", "F"), ("T", "U", "V", "W")):
        put(vv, "%s%d" % (ci, r4), "=%s" % TOT("$%s$%d" % (qc, qr)), fill=F_CALC, align=CTR)
    put(vv, "G%d" % r4, "=SUM($C%d:$F%d)" % (r4, r4), fill=F_CALC, align=CTR)
    put(vv, "H%d" % r4, '=IF($G%d=$B%d,"PASS","FAIL")' % (r4, r4), fill=F_CALC, align=CTR)
    if rk == "4급":
        put(vv, "I%d" % r4, '=IF(AND($B%d=7,$C%d=1,$D%d=2,$E%d=3,$F%d=1),"PASS — [별표 10] 7명 행(1/2/3/1)과 일치","확인 필요")'
            % (r4, r4, r4, r4, r4), font=FN_S, fill=F_CALC)
    else:
        put(vv, "I%d" % r4, "10명 이상 — [별표 10] 대상 아님", font=FN_S, fill=F_CALC)
    r4 += 1
V4_R0, V4_R1 = r4 - 2, r4 - 1
r4 += 1

sec(r4, "5. 1차·2차 평가자 간 순위 일치도 (스피어만) — 규칙 제9조⑥2호 보고 항목")
bs5, be5 = block_span[4]
bs6, be6 = block_span[5]
put(vv, "A%d" % (r4 + 1), "평가군 5(1차 4급 7명) 과 평가군 6(2차 4급 7명) 의 순위상관", font=FN, align=WRAP)
put(vv, "B%d" % (r4 + 1), "=ROUND(CORREL(%s,%s),4)"
    % (GP("$H$%d:$H$%d" % (bs5, be5)), GP("$H$%d:$H$%d" % (bs6, be6))), fill=F_CALC, align=CTR, fmt="0.0000")
put(vv, "C%d" % (r4 + 1), "순위끼리의 피어슨 상관 = 스피어만 ρ. 동점이 없으므로 이 등식이 성립한다.", font=FN_S)
r4 += 3

sec(r4, "6. 종합 판정")
put(vv, "B%d" % r4,
    '=IF(AND(COUNTIF($E$%d:$E$%d,"FAIL")=0,COUNTIF($H$%d:$H$%d,"FAIL")=0,COUNTIF($K$%d:$K$%d,"FAIL")=0,'
    'COUNTIF($F$%d:$F$%d,"FAIL")=0,COUNTIF($D$%d:$D$%d,"FAIL")=0,COUNTIF($H$%d:$H$%d,"FAIL")=0),'
    '"PASS — 전 항목 통과","FAIL — 위 항목 확인")'
    % (V1_R0, V1_R1, V1_R0, V1_R1, V1_R0, V1_R1,
       V2_R0, V2_R1, V3_R0, V3_R1, V4_R0, V4_R1),
    fill=F_CALC, align=CTR, font=Font(bold=True, size=12))
put(vv, "C%d" % r4, "① 평가군별 평균·표준편차·조정 합계  ② ÷s(n) 효과  ③ 순위 보존·절단·동점·s(n) 표  ④ 등급 정원", font=FN_S)

for rng in ("D7:D%d" % (r4 + 2), "E7:E%d" % (r4 + 2), "F7:F%d" % (r4 + 2),
            "G7:G%d" % (r4 + 2), "H7:H%d" % (r4 + 2), "I7:I%d" % (r4 + 2), "J7:J%d" % (r4 + 2),
            "K7:K%d" % (r4 + 2),
            "B7:B%d" % (r4 + 2)):
    vv.conditional_formatting.add(rng, CellIsRule(operator="containsText",
        formula=['NOT(ISERROR(SEARCH("PASS",%s)))' % rng.split(":")[0]],
        fill=PatternFill("solid", fgColor="C6EFCE")))
    vv.conditional_formatting.add(rng, CellIsRule(operator="containsText",
        formula=['NOT(ISERROR(SEARCH("FAIL",%s)))' % rng.split(":")[0]],
        fill=PatternFill("solid", fgColor="FFC7CE")))
widths(vv, {"A": 46, "B": 22, "C": 18, "D": 18, "E": 14, "F": 20, "G": 20, "H": 14,
            "I": 22, "J": 16, "K": 12})
vv.freeze_panes = "B5"

# ════════════════════════════════════════════════════════════════════
#  ⑥ 예제
# ════════════════════════════════════════════════════════════════════
ex = wb.create_sheet(SH_EX)
put(ex, "A1", "⑥ 예제 — 세 가지 상황의 재현", font=B_TTL)
put(ex, "A2", DISCLAIMER, fill=F_NOTE, font=FN)
put(ex, "A3", "각 예제는 다른 시트를 참조하지 않고 자기 안에서 수식으로 완결된다(④ s(n) 표만 참조).", font=FN_S)

def example_block(row, title, oneline, raws, adj=None, base_ref=None, note=None):
    """원점수 → 순위 → z → 보정점수 를 수식으로 재현하는 블록."""
    put(ex, "A%d" % row, title, font=B_SEC)
    put(ex, "A%d" % (row + 1), "무엇을 보여주는 예제인가 — " + oneline, font=FN, fill=F_NOTE, align=WRAP)
    ex.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=9)
    h = row + 2
    cols = ["피평가자", "원점수", "인원 n", "순위 r\n(낮은 순=1위)", "s(n)", "p", "z = Φ⁻¹(p)",
            "보정점수\n÷s(n) 적용 (별표 9 최종본)", "참고 · ÷s(n) 미적용\n(세부지침 §1-2 표기값)"]
    if adj is not None:
        cols += ["조정 (제7호)", "조정 후 점수"]
    hdr(ex, h, cols)
    ex.row_dimensions[h].height = 40
    d0 = h + 1
    d1 = d0 + len(raws) - 1
    for k, v in enumerate(raws):
        rr = d0 + k
        put(ex, "A%d" % rr, "직원 %s" % chr(ord("A") + k), fill=F_CALC, align=CTR)
        put(ex, "B%d" % rr, v, fill=F_IN, align=CTR, fmt="0.0##")
        put(ex, "C%d" % rr, "=COUNT($B$%d:$B$%d)" % (d0, d1), fill=F_CALC, align=CTR)
        put(ex, "D%d" % rr, "=RANK.EQ($B%d,$B$%d:$B$%d,1)" % (rr, d0, d1), fill=F_CALC, align=CTR)
        put(ex, "E%d" % rr, "=IF($C%d>=17,1,INDEX(%s,MATCH($C%d,%s,0)))" % (rr, SN_APPLY, rr, SN_NCOL),
            fill=F_CALC, align=CTR, fmt="0.0000")
        put(ex, "F%d" % rr, "=($D%d-0.375)/($C%d+0.25)" % (rr, rr), fill=F_CALC, align=CTR, fmt="0.0000")
        put(ex, "G%d" % rr, "=NORM.S.INV($F%d)" % rr, fill=F_CALC, align=CTR, fmt="0.0000")
        put(ex, "H%d" % rr, "=ROUND(MEDIAN(60,80+7*$G%d/$E%d,100),3)" % (rr, rr),
            fill=F_CALC, align=CTR, fmt="0.000")
        put(ex, "I%d" % rr, "=ROUND(MEDIAN(60,80+7*$G%d,100),3)" % rr, fill=F_CALC, align=CTR, fmt="0.000")
        if adj is not None:
            put(ex, "J%d" % rr, adj[k], fill=F_IN, align=CTR)
            put(ex, "K%d" % rr, "=ROUND(MEDIAN(60,$H%d+$J%d,100),3)" % (rr, rr),
                fill=F_CALC, align=CTR, fmt="0.000")
    s = d1 + 1
    put(ex, "A%d" % s, "평균", font=B_HEAD, fill=F_HEAD, align=CTR)
    put(ex, "B%d" % s, "=ROUND(AVERAGE($B$%d:$B$%d),3)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "H%d" % s, "=ROUND(AVERAGE($H$%d:$H$%d),3)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "I%d" % s, "=ROUND(AVERAGE($I$%d:$I$%d),3)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "A%d" % (s + 1), "표준편차 STDEVP", font=B_HEAD, fill=F_HEAD, align=CTR)
    put(ex, "B%d" % (s + 1), "=ROUND(STDEVP($B$%d:$B$%d),3)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "H%d" % (s + 1), "=ROUND(STDEVP($H$%d:$H$%d),3)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "I%d" % (s + 1), "=ROUND(STDEVP($I$%d:$I$%d),3)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "A%d" % (s + 2), "부여폭 (최고−최저)", font=B_HEAD, fill=F_HEAD, align=CTR)
    put(ex, "B%d" % (s + 2), "=ROUND(MAX($B$%d:$B$%d)-MIN($B$%d:$B$%d),3)" % (d0, d1, d0, d1),
        fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "H%d" % (s + 2), "=ROUND(MAX($H$%d:$H$%d)-MIN($H$%d:$H$%d),3)" % (d0, d1, d0, d1),
        fill=F_CALC, align=CTR, fmt="0.000")
    nxt = s + 3
    if adj is not None:
        put(ex, "J%d" % s, "=ROUND(SUM($J$%d:$J$%d),2)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.00")
        put(ex, "K%d" % s, "=ROUND(AVERAGE($K$%d:$K$%d),3)" % (d0, d1), fill=F_CALC, align=CTR, fmt="0.000")
        put(ex, "L%d" % s, '=IF($J%d=0,"PASS — 조정 합계 0 (제7호)","FAIL")' % s, fill=F_CALC, align=CTR)
        ex.conditional_formatting.add("L%d" % s, CellIsRule(operator="containsText",
            formula=['NOT(ISERROR(SEARCH("PASS",L%d)))' % s], fill=PatternFill("solid", fgColor="C6EFCE")))
    if note:
        put(ex, "A%d" % nxt, note, font=FN_S, align=WRAP)
        ex.merge_cells(start_row=nxt, start_column=1, end_row=nxt, end_column=11)
        nxt += 1
    return nxt + 1, d0, d1

nr, A_d0, A_d1 = example_block(
    5, "예제 A · 10명 평가군 — 극단 성향 평가자 (부여폭 넓음)",
    "부여폭 38점을 보정하면 표준편차 7.00점으로 압축된다. 부서 내 순위는 그대로 보존된다.",
    G1_RAW, adj=G1_ADJ,
    note="H열이 개정안 [별표 9] 최종본(÷s(n) 적용)이다. I열은 ÷s(n)이 없던 단계의 값으로, "
         "〈08 세부지침〉 §1-2 표(90.83 … 69.17)와 같다. 최종본에서는 92.137 … 67.863 이 되고 "
         "표준편차가 6.244 → 7.000 으로 제4호의 목표값을 달성한다.")

nr2, B_d0, B_d1 = example_block(
    nr, "예제 B · 4명 평가군 — [별표 9] 제5호 적용 (1차 보정 미실시 · 2차 100%)",
    "6명 미만이므로 제3호 보정을 실시하지 아니하고, 2차 평가자의 보정점수를 100% 반영한다. 1차 평가자의 ±5점 조정만 남는다.",
    G2_RAW, adj=None,
    note="아래 표가 제5호의 실제 계산이다. 위 H·I열의 값은 「4명 평가군에 제3호를 그대로 적용하면 어떻게 되는가」를 "
         "참고로 보여주는 것이며, 제5호에 따라 사용하지 아니한다.")

# 예제 B — 제5호 실계산
put(ex, "A%d" % nr2, "예제 B 계산 — 제5호에 따른 실제 산출", font=B_SEC)
hdr(ex, nr2 + 1, ["피평가자", "1차 원점수\n(참고)", "2차 평가군 인원", "2차 평가군 순위",
                  "2차 s(n)", "2차 보정점수\n= 업적·역량 100%", "1차 조정 (제7호)", "최종 점수"])
ex.row_dimensions[nr2 + 1].height = 34
B2_N, B2_RANKS = 36, [6, 15, 24, 31]     # 세부지침 §1-3 과 동일한 조건
b0 = nr2 + 2
for k in range(4):
    rr = b0 + k
    put(ex, "A%d" % rr, "직원 %s" % chr(ord("A") + k), fill=F_CALC, align=CTR)
    put(ex, "B%d" % rr, G2_RAW[k], fill=F_IN, align=CTR)
    put(ex, "C%d" % rr, B2_N, fill=F_IN, align=CTR)
    put(ex, "D%d" % rr, B2_N - B2_RANKS[k] + 1, fill=F_IN, align=CTR)
    put(ex, "E%d" % rr, "=IF($C%d>=17,1,INDEX(%s,MATCH($C%d,%s,0)))" % (rr, SN_APPLY, rr, SN_NCOL),
        fill=F_CALC, align=CTR, fmt="0.0000")
    put(ex, "F%d" % rr, "=ROUND(MEDIAN(60,80+7*NORM.S.INV(($D%d-0.375)/($C%d+0.25))/$E%d,100),3)"
        % (rr, rr, rr), fill=F_CALC, align=CTR, fmt="0.000")
    put(ex, "G%d" % rr, G2_ADJ[k], fill=F_IN, align=CTR)
    put(ex, "H%d" % rr, "=ROUND(MEDIAN(60,$F%d+$G%d,100),3)" % (rr, rr), fill=F_CALC, align=CTR, fmt="0.000")
b1 = b0 + 3
put(ex, "A%d" % (b1 + 1), "합계 / 평균", font=B_HEAD, fill=F_HEAD, align=CTR)
put(ex, "F%d" % (b1 + 1), "=ROUND(AVERAGE($F$%d:$F$%d),3)" % (b0, b1), fill=F_CALC, align=CTR, fmt="0.000")
put(ex, "G%d" % (b1 + 1), "=ROUND(SUM($G$%d:$G$%d),2)" % (b0, b1), fill=F_CALC, align=CTR, fmt="0.00")
put(ex, "H%d" % (b1 + 1), "=ROUND(AVERAGE($H$%d:$H$%d),3)" % (b0, b1), fill=F_CALC, align=CTR, fmt="0.000")
put(ex, "I%d" % (b1 + 1), '=IF(AND($G%d=0,ROUND($F%d,3)=ROUND($H%d,3)),"PASS — 조정 합계 0이므로 평가군 평균 불변","FAIL")'
    % (b1 + 1, b1 + 1, b1 + 1), fill=F_CALC, align=CTR)
ex.conditional_formatting.add("I%d" % (b1 + 1), CellIsRule(operator="containsText",
    formula=['NOT(ISERROR(SEARCH("PASS",I%d)))' % (b1 + 1)], fill=PatternFill("solid", fgColor="C6EFCE")))
put(ex, "A%d" % (b1 + 2), "2차 평가자의 평가군 인원도 6명 미만이면 80점을 부여한다([별표 9] 제5호 후단). "
                          "위 조건(2차 36명 · 순위 6/15/24/31)은 〈08 세부지침〉 §1-3 예제와 같고, "
                          "n=36 ≥ 17 이므로 s(n)=1 이어서 세부지침의 87.10 / 81.71 / 77.27 / 72.90 과 일치한다.", font=FN_S)
nr3 = b1 + 5

nr4, C_d0, C_d1 = example_block(
    nr3, "예제 C · 10명 평가군 — 균등 성향 평가자 (전원 78~80점)  ★ 개정안의 핵심 효과",
    "부여폭 1.8점밖에 없던 평가군에 보정 후 표준편차 7.00점의 차등이 생긴다. 예제 A와 보정점수가 완전히 같아진다.",
    [round(78.0 + 0.2 * k, 1) for k in range(10)], adj=None,
    note="예제 A(부여폭 38점)와 예제 C(부여폭 1.8점)의 H열 보정점수가 소수점까지 같다. "
         "이것이 [별표 9]가 하는 일이다 — 평가자가 넓게 주든 좁게 주든 결과가 같아지고, "
         "차등이 없던 평가군에 차등이 생긴다.")
# 예제 C의 각 행에 「같은 순위인 예제 A 행과의 차이」를 넣어 행별로 대조한다.
put(ex, "M%d" % (C_d0 - 1), "예제 A 같은 순위와의 차", fill=F_HEAD, font=B_HEAD, align=CTR)
for rr in range(C_d0, C_d1 + 1):
    put(ex, "M%d" % rr,
        '=IFERROR(ROUND($H%d,3)-ROUND(INDEX($H$%d:$H$%d,MATCH($D%d,$D$%d:$D$%d,0)),3),"예제 A에 같은 순위가 없다")'
        % (rr, A_d0, A_d1, rr, A_d0, A_d1), fill=F_CALC, align=CTR, fmt="0.000")
put(ex, "A%d" % nr4, "예제 A와 예제 C의 보정점수 동일성 검증", font=B_SEC)
put(ex, "B%d" % nr4, '=IF(COUNTIF($M$%d:$M$%d,"<>0")=0,'
                     '"PASS — 두 예제의 보정점수가 전부 일치한다 (차이 0)","FAIL — M열 확인")'
    % (C_d0, C_d1), fill=F_CALC, align=CTR, font=Font(bold=True, size=11))
put(ex, "H%d" % nr4, "차이 합계", font=B_HEAD, fill=F_HEAD, align=CTR)
put(ex, "I%d" % nr4, '=IFERROR(ROUND(SUM($M$%d:$M$%d),3),"산출 불가 — M열 확인")' % (C_d0, C_d1),
    fill=F_CALC, align=CTR, fmt="0.000")
ex.conditional_formatting.add("B%d" % nr4, CellIsRule(operator="containsText",
    formula=['NOT(ISERROR(SEARCH("PASS",B%d)))' % nr4], fill=PatternFill("solid", fgColor="C6EFCE")))
ex.conditional_formatting.add("B%d" % nr4, CellIsRule(operator="containsText",
    formula=['NOT(ISERROR(SEARCH("FAIL",B%d)))' % nr4], fill=PatternFill("solid", fgColor="FFC7CE")))
widths(ex, {"A": 20, "B": 14, "C": 16, "D": 16, "E": 12, "F": 20, "G": 14,
            "H": 26, "I": 26, "J": 12, "K": 14, "L": 40, "M": 22})
ex.freeze_panes = "B5"

# ════════════════════════════════════════════════════════════════════
#  ⑦ MIS 요구사항
# ════════════════════════════════════════════════════════════════════
ms = wb.create_sheet(SH_MIS)
put(ms, "A1", "⑦ MIS 요구사항 — 내부경영정보시스템 근무평정 모듈 기능 명세(요청안)", font=B_TTL)
put(ms, "A2", "이 계산서의 계산을 시스템에 넣기 위한 기능 요건이다. 우선순위 ★는 규정 시행에 반드시 필요한 항목이다.", font=FN_S)
hdr(ms, 4, ["구분", "No", "요구 기능", "상세 요건", "근거", "우선순위"])
MIS = [
 ("화면", "평가자 점수 입력 화면", "피평가자별 원점수(60~100)를 입력한다. 평가군 인원 n과 입력 진행률을 함께 표시한다.", "규칙 제9조① · [별표 1]", "★"),
 ("화면", "동점 입력 차단", "같은 평가군 내에서 이미 사용된 점수를 다시 입력하면 저장을 차단하고, 사용 가능한 점수를 안내한다. 경고가 아니라 저장 차단이어야 한다 — 동점이 있으면 순위가 중복되어 [별표 9] 제2호·제3호가 성립하지 않는다.", "개정 2 (제9조⑥) · [별표 9] 제2호", "★"),
 ("화면", "보정 결과 비공개", "평가자 화면에는 원점수만 표시한다. 보정점수·보정 후 순위를 평가자에게 미리 보여주지 아니한다.", "〈08 세부지침〉 §4-2", "★"),
 ("화면", "제7호 조정 화면", "1차 평가자에게 평가군 단위로 보정점수 목록과 조정 입력란(±5점)을 제시하고, 조정 합계를 실시간으로 표시한다.", "[별표 9] 제7호", "★"),
 ("입력", "조정 합계 0 저장 차단", "평가군 내 조정점수 합계가 0이 아니면 저장을 차단한다. 잔차가 남은 경우 제7호 가·나·다 순서의 배분 결과를 제시하고, 배분 후에도 남는 잔차는 별도 항목으로 기록한다.", "[별표 9] 제7호 가·나·다", "★"),
 ("입력", "조정 사유 필수 입력", "조정량이 0이 아닌 피평가자는 사유 입력 없이 저장할 수 없다. 사유는 평가표에 함께 저장한다.", "[별표 9] 제7호 후단", "★"),
 ("입력", "±5점 범위 검증", "개별 조정량이 −5 ~ +5 를 벗어나면 입력을 거부한다.", "[별표 9] 제7호", "★"),
 ("입력", "다면평가 평가자 지정 이력", "피평가자 추천(5인 이상) → 부서장 확인 → 인사담당부서장 지정의 3단계 이력을 각 단계 일시와 처리자와 함께 보관한다. 1차 평가자는 다면평가자로 지정할 수 없도록 시스템이 차단한다.", "개정 15 (제22조①②③)", "★"),
 ("계산", "보정 일괄 실행", "평가 마감 후 인사담당부서만 보정을 실행한다. 평가자별·평가구분별·직급군별로 평가군을 자동 구성하고 [별표 9] 제3호를 적용한다. 실행 시점·실행자를 로그로 남긴다.", "제9조④ · 〈08 세부지침〉 §4-2", "★"),
 ("계산", "s(n) 내장", "n = 2~40 의 s(n)을 계산식(정규 순서통계량의 표준편차)으로 내장하고 n ≥ 17 은 1을 적용한다. 상수표를 하드코딩하지 아니한다. 제3호 후단이 개정되어 s(n) 적용 구간이 확대되는 경우 설정 변경만으로 대응할 수 있게 설계한다.", "[별표 9] 제3호", "★"),
 ("계산", "제5호 자동 판정", "평가군 인원이 6명 미만이면 1차 보정을 건너뛰고 2차 보정점수를 100% 반영한다. 2차 평가군도 6명 미만이면 80점을 부여한다. 적용 평가군 수를 집계한다.", "[별표 9] 제5호 · 제9조⑦3호", "★"),
 ("계산", "소수점 셋째자리 반올림", "보정점수·조정 후 점수·종합점수를 소수점 셋째자리에서 반올림하여 저장한다. 표시 단계가 아니라 저장 단계에서 확정한다.", "[별표 9] 제3호", "★"),
 ("계산", "절단 처리", "60점 미만은 60점, 100점 초과는 100점으로 절단하고 절단 발생 건수를 집계한다.", "[별표 9] 제3호", "○"),
 ("계산", "다면평가 보정", "유효 응답 6인 이상이면 최고점·최저점을 각각 하나씩 제외한 뒤 기관 전체를 하나의 평가군으로 하여 제3호를 적용한다. 유효 응답 3인 미만은 80점으로 하고 업적·역량 배점을 그 비율만큼 가산한다.", "개정 15 (제22조⑤⑦)", "★"),
 ("계산", "등급 배정 (직급별·최대잔여법)", "직급별 전사 단위로 S20·A30·B40·C10%를 최대잔여법(정수부 우선 배정 → 소수부 큰 순서로 잔여 배정)으로 배정한다. 부서 단위 배정을 허용하지 아니한다. 10명 미만은 [별표 10] 표를 적용한다.", "규칙 제26조①③ · [별표 10]", "★"),
 ("계산", "평정점 산출", "등급 구간(S100·A90·B80·C70) 안에서 「구간 최고점 − (10 ÷ 등급별 인원) × (등급 내 순위 − 1)」로 산출한다.", "[별표 7] · 제27조④", "★"),
 ("계산", "배점 전환", "업적·역량 : 다면 : 공통 배점을 설정값으로 관리한다(개정 14 채택 시 75:20:5, 현행 85:10:5). 부서장은 1차:2차 = 50:50 및 상향평가를 적용한다.", "제9조② · 개정 14 · [별표 2]", "○"),
 ("검증", "평가군별 보정 전·후 평균·표준편차", "평가군별로 보정 전·후의 평균과 표준편차를 자동 산출한다. 보정 후 평균 80.00 · 표준편차 7.00 을 만족하지 않으면 오류로 표시한다.", "제9조⑥1호 (보고 항목)", "★"),
 ("검증", "1차·2차 평가자 간 순위 일치도 (스피어만)", "같은 피평가자 집단에 대한 1차 평가자 순위와 2차 평가자 순위의 스피어만 순위상관을 평가군별·직급별로 자동 산출한다.", "제9조⑥2호", "★"),
 ("검증", "직급별 순위 변동 건수", "보정 전 종합점수 순위와 보정 후 순위를 비교해 직급별 변동 건수를 산출한다.", "제9조⑦2호", "★"),
 ("검증", "부서 내 순위 보존 확인", "평가군별로 원점수 순위와 보정점수 순위가 일치하는지 전건 확인하고 불일치 건을 오류로 표시한다(동점 입력의 사후 탐지).", "[별표 9] 제2호·제3호", "○"),
 ("검증", "제6호 병산", "시행 첫해에는 제3호와 평균·표준편차 일치법을 함께 산출해 비교표를 출력한다.", "[별표 9] 제6호", "○"),
 ("출력", "보정 전·후 점수 동시 보관", "원점수 · 보정점수 · 조정량 · 조정 후 점수 · 종합점수 · 평정점을 각각 별도 항목으로 영구 보관한다. 덮어쓰지 아니한다.", "제28조① (공개 대상) · [별표 9]", "★"),
 ("출력", "본인 공개 화면", "피평가자에게 보정 전·후의 점수와 평정점을 공개한다. 1·2차 평가자별 점수와 직급 내 순위는 공개하지 아니한다.", "제28조①", "★"),
 ("출력", "노사협의회 보고서 자동 생성", "제9조⑥·⑦의 보고 항목(평가군별 보정 전·후 평균·표준편차, 직급별 순위 변동 건수, 제5호 적용 평가군 수, 다면평가 평가자별 부여점수 분포)을 한 장으로 출력한다.", "제9조⑥⑦ · 제22조⑧", "★"),
 ("출력", "이의신청 근거 자료", "본인 요청 시 자기 평가군의 인원 n, 자기 순위 r, 적용 s(n), 계산 과정을 제시한다.", "제29조 · [별표 9] 제3호", "○"),
 ("권한", "역할 분리", "평가자(원점수 입력) · 1차 평가자(제7호 조정) · 인사담당부서(보정 실행·등급 배정) · 감사(열람) 권한을 분리한다. 인사담당부서는 원점수를 수정할 수 없다.", "제9조④ · 제28조", "★"),
 ("권한", "변경 이력", "모든 점수 항목의 변경에 대해 변경 전·후 값, 변경자, 변경 일시, 사유를 남긴다.", "제28조 · 〈08 세부지침〉 §4-5", "★"),
 ("권한", "다면평가 익명성", "다면평가 평가자 명단과 개별 부여점수는 인사담당부서와 감사 외에 조회할 수 없도록 한다. 분포 집계는 익명으로만 출력한다.", "개정 15 (제22조④⑧)", "★"),
]
mr = 5
for grp, fn_, det, law, pri in MIS:
    put(ms, "A%d" % mr, grp, fill=F_CALC, align=CTR, font=B_HEAD)
    put(ms, "B%d" % mr, mr - 4, fill=F_CALC, align=CTR)
    put(ms, "C%d" % mr, fn_, fill=F_CALC, align=WRAP, font=Font(bold=True, size=10))
    put(ms, "D%d" % mr, det, fill=F_CALC, align=WRAP)
    put(ms, "E%d" % mr, law, fill=F_CALC, align=WRAP, font=FN_S)
    put(ms, "F%d" % mr, pri, fill=F_CALC, align=CTR)
    ms.row_dimensions[mr].height = 46
    mr += 1
widths(ms, {"A": 8, "B": 5, "C": 32, "D": 84, "E": 26, "F": 8})
ms.freeze_panes = "A5"

# ════════════════════════════════════════════════════════════════════
#  ⑧ 근거 조문
# ════════════════════════════════════════════════════════════════════
lw = wb.create_sheet(SH_LAW)
put(lw, "A1", "⑧ 근거 조문 — 이 계산서의 각 계산이 개정안의 어느 조문에 대응하는가", font=B_TTL)
put(lw, "A2", "정본은 docs/amendments.md 이다. 이 표는 대응관계만 정리한 것이다.", font=FN_S)
hdr(lw, 4, ["시트 · 항목", "계산 내용", "근거 조문", "비고"])
LAW = [
 ("② 평가군 계산 A~D열", "평가자 · 평가구분 · 직급군으로 평가군을 구성한다. 1차와 2차를 각각 별도로 보정한다.", "[별표 9] 제1호", "1차·2차 반영비율은 [별표 2]"),
 ("② H열 순위 r", "평가군 내에서 낮은 점수부터 1위를 부여한다(최저점자 1위 · 최고점자 n위).", "[별표 9] 제2호", "RANK.EQ(…,1) 오름차순"),
 ("② Q열 동점 경고", "동일 평가군 내 동점 입력을 탐지한다.", "개정 2 (제9조⑥) · [별표 9] 제2호", "동점은 순위 중복을 낳아 산식을 깨뜨린다"),
 ("② I열 s(n)", "인원 n일 때 z값들의 표준편차. 17명 이상은 1로 본다.", "[별표 9] 제3호", "④ s(n) 표에서 계산"),
 ("② J~L열 보정점수", "z = Φ⁻¹((r−0.375)/(n+0.25)) · 보정점수 = 80 + 7 × z ÷ s(n)", "[별표 9] 제3호", "Blom 정규 순서통계량"),
 ("② L열 절단·반올림", "60점 미만은 60점, 100점 초과는 100점. 소수점 셋째자리 반올림.", "[별표 9] 제3호", "MEDIAN(60,x,100) · ROUND(x,3)"),
 ("기준값 80 · 7", "80은 [별표 7] 평정점 구간(60~100)의 중앙값, 7은 정수로서 검산 재현성이 보장되는 값.", "[별표 9] 제4호", "노사 협의로 달리 정할 수 있다"),
 ("② M열 제5호", "인원 6명 미만이면 제3호 보정을 실시하지 아니하고 2차 보정점수를 100% 반영한다. 2차도 6명 미만이면 80점.", "[별표 9] 제5호", "적용 평가군 수는 제9조⑦3호 보고 항목"),
 ("② N~P열 제7호", "1차 평가자가 ±5점 범위에서 조정한다. 평가군 내 합계는 0. 조정 사유를 기재한다.", "[별표 9] 제7호", "잔차 배분 순서는 같은 호 가·나·다"),
 ("② R~S열 순위 보존", "원점수 순위와 보정점수 순위가 일치하는지 확인한다.", "[별표 9] 제2호·제3호", "단조변환이므로 정의상 100% 보존"),
 ("③ H~J열 업적·역량", "1차 60% + 2차 40%. 부서장(4급 이상)은 50:50.", "[별표 2] · 규칙 제9조②", "제5호 적용 시 2차 100%"),
 ("③ K~P열 다면 보정", "유효 응답 6인 이상이면 최고·최저점을 각각 하나씩 제외한 후 기관 전체를 하나의 평가군으로 하여 제3호를 적용한다. 3인 미만은 80점.", "개정 15 (제22조⑤⑦) · [별표 9]", "부서장은 상향평가(제22조의2②)"),
 ("③ B3 배점 스위치", "개정 14 채택 시 업적·역량 75 : 다면 20 : 공통 5. 미채택 시 85 : 10 : 5.", "규칙 제9조② · 개정 14 · 개정 5", "개정 15 미채택 시 개정 14도 철회"),
 ("③ T열 종합점수", "업적·역량 배점 × 업적·역량 점수 + 다면 배점 × 다면 보정점수 + 5% × 공통평가", "규칙 제9조① · 제3조8호(개정 1)", "가감평가점수는 이 계산서 범위 외"),
 ("③ U~Z열 정원", "직급별 전사 단위로 S20 · A30 · B40 · C10%를 최대잔여법으로 배정한다.", "규칙 제26조①③", "부서 단위 배정이 아니다"),
 ("③ 정원 블록", "① 인원×비율의 정수부 배정 ② 남은 자리는 소수부가 큰 등급 순으로 하나씩", "규칙 제26조②③ · 〈09 수치의 근거〉 §2-3", "10명 미만은 [별표 10]"),
 ("③ AD열 평정점", "구간 최고점 − (10 ÷ 등급별 인원) × (등급 내 순위 − 1)", "[별표 7] · 규칙 제27조④", "S100 · A90 · B80 · C70"),
 ("④ s(n) 표", "n = 2~40 의 z값과 그 표준편차를 NORM.S.INV · STDEVP 로 계산한다.", "[별표 9] 제3호", "n=2~16 은 개정안 표 참고값과 대조"),
 ("④·⑤ n ≥ 17 구간", "제3호 후단에 따라 s(n) = 1 을 적용한다. 그 결과 17명 이상 평가군의 표준편차는 7.00에 미달한다(n=17 → 6.51 · n=30 → 6.70 · n=40 → 6.76). 1차 평가군(부서 단위 3~16명)은 영향이 없고 2차 평가군(본부 단위)이 영향을 받는다.", "[별표 9] 제3호 후단", "★ 노사 협의 논점. ⑤ 검증 1-2절 참조. s(n) 표를 인원 구간 전체로 확대하면 해소된다"),
 ("⑤ 검증 1·2", "평가군별 평균 80.00 · 표준편차 7.00 을 판정하고, ÷s(n) 없을 때의 미달값을 함께 보인다.", "[별표 9] 제3호·제4호", "제9조⑥1호 보고 항목과 같은 통계"),
 ("⑤ 검증 5", "1차·2차 평가자 간 순위 일치도(스피어만)", "규칙 제9조⑥2호", "순위끼리의 피어슨 상관"),
 ("⑦ MIS 요구사항", "보정 실행 주체·시점, 보정 전·후 점수 보관, 보고 항목 자동 산출", "제9조④⑥⑦ · 제28조① · 제22조⑧", "제31조의2에 따라 노사협의회 협의 대상"),
]
lr = 5
for a, b, c, d in LAW:
    put(lw, "A%d" % lr, a, fill=F_CALC, align=WRAP, font=Font(bold=True, size=10))
    put(lw, "B%d" % lr, b, fill=F_CALC, align=WRAP)
    put(lw, "C%d" % lr, c, fill=F_CALC, align=WRAP)
    put(lw, "D%d" % lr, d, fill=F_CALC, align=WRAP, font=FN_S)
    lw.row_dimensions[lr].height = 40
    lr += 1
widths(lw, {"A": 26, "B": 80, "C": 34, "D": 34})
lw.freeze_panes = "A5"

# ════════════════════════════════════════════════════════════════════
#  ① 사용법  — 맨 앞으로 이동
# ════════════════════════════════════════════════════════════════════
us = wb.create_sheet(SH_USE, 0)
put(us, "A1", "창업진흥원 근무평정 점수 보정 실무 계산서", font=Font(bold=True, size=16))
put(us, "A2", "개정안 [별표 9] 「평가자별 평가점수의 보정 방법」을 엑셀 수식으로 구현한 파일", font=B_SEC)
put(us, "A4", DISCLAIMER, fill=F_NOTE, font=Font(bold=True, size=11), align=WRAP)
us.merge_cells("A4:F4"); us.row_dimensions[4].height = 32
put(us, "A5", "이 파일은 조합이 2026년 2분기 노사협의회에 제출하는 개정안 [별표 9]의 구현이다. "
              "정본은 docs/amendments.md 이며, 이 파일과 정본이 어긋나면 정본이 우선한다.", font=FN, align=WRAP)
us.merge_cells("A5:F5"); us.row_dimensions[5].height = 30

put(us, "A7", "이 파일을 만든 이유", font=B_SEC)
for k, t in enumerate([
    "1. 검산 가능성의 증명 — 조합은 「검산할 수 없는 산식은 규정에 담을 수 없다」를 원칙으로 세웠고, "
    "그 이유로 혼합효과 모형·Bradley-Terry·경험베이즈 등 고급 기법을 전부 폐기했다. "
    "이 파일은 개정안 산식이 엑셀 수식 몇 줄로 구현된다는 물증이다. 모든 셀은 값이 아니라 수식이다 — 셀을 클릭해 확인하라.",
    "2. MIS 개선 요청 명세 — 내부경영정보시스템 근무평정 모듈에 이 계산을 넣도록 요청할 때의 기능 명세서가 된다(⑦ 시트).",
]):
    put(us, "A%d" % (8 + k), t, font=FN, align=WRAP)
    us.merge_cells(start_row=8 + k, start_column=1, end_row=8 + k, end_column=6)
    us.row_dimensions[8 + k].height = 46

put(us, "A11", "셀 색 구분", font=B_SEC)
put(us, "A12", "  입력 셀 — 실무자가 채운다", fill=F_IN, font=Font(bold=True, size=10))
put(us, "B12", "원점수 · 제7호 조정량 · 조정 사유 · 다면평가 점수 · 유효 응답 수 · 공통평가 점수 · 부서장 여부 · 배점 스위치", font=FN, align=WRAP)
put(us, "A13", "  자동 계산 셀 — 손대지 않는다", fill=F_CALC, font=Font(bold=True, size=10))
put(us, "B13", "인원 n · 순위 r · s(n) · z · 보정점수 · 조정 후 점수 · 종합점수 · 등급 · 평정점 · 모든 검증 판정", font=FN, align=WRAP)
put(us, "A14", "  고지 · 경고", fill=F_NOTE, font=Font(bold=True, size=10))
put(us, "B14", "동점 입력 경고, 모의 데이터 고지 등 반드시 읽어야 하는 문구", font=FN)
put(us, "A15", "  PASS / FAIL", fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(bold=True, size=10))
put(us, "B15", "⑤ 검증 시트의 판정. 녹색이 PASS, 적색이 FAIL이다. FAIL이 하나라도 있으면 입력을 다시 확인하라.", font=FN, align=WRAP)

put(us, "A17", "작업 순서", font=B_SEC)
STEPS = [
 ("1", "② 평가군 계산", "평가자명 · 평가군 구분(1차/2차) · 직급 · 피평가자 · 원점수를 입력한다. "
  "평가군은 「동일 평가자 · 동일 평가구분 · 동일 직급군」이다([별표 9] 제1호). 평가군마다 행을 붙여서 쓴다."),
 ("2", "② 동점 확인", "4행의 「동점 입력 건수」가 0건인지 확인한다. 동점이 있으면 순위가 중복되어 산식이 깨진다. "
  "개정 2가 동점을 금지하므로 정상 운용에서는 발생하지 않지만 입력 오류를 여기서 잡는다."),
 ("3", "② 제5호 확인", "인원 n이 6명 미만인 평가군은 「제5호 미실시」로 자동 표시되고, "
  "2차 평가자의 보정점수가 자동으로 기준점수가 된다. 2차도 6명 미만이면 80점이 들어간다."),
 ("4", "② 제7호 조정", "1차 평가자가 조정란(±5점)에 입력한다. ⑤ 검증 시트에서 평가군별 합계가 0.00인지 확인한다. "
  "조정한 피평가자는 사유란을 반드시 채운다. 2차 평가자는 조정하지 아니한다."),
 ("5", "③ 종합점수·등급", "다면(상향)평가 점수 · 유효 응답 수 · 공통평가 점수 · 부서장 여부를 입력한다. "
  "B3의 배점 스위치로 개정 14 채택안(75:20:5)과 현행(85:10:5)을 전환한다."),
 ("6", "③ 등급 확인", "등급 정원은 직급별 전사 단위로 최대잔여법에 따라 자동 배정된다(시트 하단 정원 블록). "
  "부서 단위 배정이 아니다(규칙 제26조①)."),
 ("7", "⑤ 검증", "모든 판정이 PASS인지 확인한다. 종합 판정 셀 하나만 봐도 된다."),
 ("8", "인쇄 · 제출", "사측 설명 시에는 ⑤ 검증 → ⑥ 예제 → ④ s(n) 표 순서로 보인다. "
  "④는 「그 표는 어디서 나왔는가」에 대한 답이다."),
]
hdr(us, 18, ["순서", "시트", "무엇을 하는가"])
ur = 19
for a, b, c in STEPS:
    put(us, "A%d" % ur, a, fill=F_CALC, align=CTR, font=B_HEAD)
    put(us, "B%d" % ur, b, fill=F_CALC, align=CTR, font=Font(bold=True, size=10))
    put(us, "C%d" % ur, c, fill=F_CALC, align=WRAP)
    us.merge_cells(start_row=ur, start_column=3, end_row=ur, end_column=6)
    us.row_dimensions[ur].height = 44
    ur += 1

ur += 1
put(us, "A%d" % ur, "시트 구성", font=B_SEC); ur += 1
hdr(us, ur, ["시트", "내용"]); ur += 1
for a, b in [
 (SH_USE, "이 시트. 색 구분과 작업 순서."),
 (SH_GRP, "★ 주 시트. 평가군 구성 · 순위 · s(n) · z · 보정점수 · 제7호 조정. 실무자는 여기서만 작업한다."),
 (SH_TOT, "1차·2차 보정점수와 다면·공통을 합산한 종합점수, 직급별 등급(최대잔여법)과 평정점([별표 7])."),
 (SH_SN, "n = 2~40 의 z값과 그 표준편차를 NORM.S.INV · STDEVP 로 계산한다. 「s(n) 표는 어디서 나왔는가」에 대한 답."),
 (SH_VAL, "★ 사측 설득의 핵심. 평균 80·표준편차 7 달성, ÷s(n) 없을 때의 미달, 조정 합계 0, 순위 보존, 절단 0건, 정원 합계를 PASS/FAIL로 자동 판정한다."),
 (SH_EX, "예제 A(극단 성향) · B(제5호) · C(균등 성향)를 수식으로 재현한다."),
 (SH_MIS, "내부경영정보시스템 근무평정 모듈에 요청할 기능 명세(화면·입력·계산·검증·출력·권한)."),
 (SH_LAW, "각 계산이 개정안의 어느 조문에 대응하는지 정리한 표."),
]:
    put(us, "A%d" % ur, a, fill=F_CALC, align=CTR, font=Font(bold=True, size=10))
    put(us, "B%d" % ur, b, fill=F_CALC, align=WRAP)
    us.merge_cells(start_row=ur, start_column=2, end_row=ur, end_column=6)
    us.row_dimensions[ur].height = 40
    ur += 1

ur += 1
put(us, "A%d" % ur, "산식 요약 — [별표 9] 제3호", font=B_SEC); ur += 1
put(us, "A%d" % ur, "z = Φ⁻¹( (r − 0.375) ÷ (n + 0.25) )        보정점수 = 80 + 7 × z ÷ s(n)",
    font=Font(bold=True, size=12), fill=F_CALC); ur += 1
put(us, "A%d" % ur, "r = 평가군 내 순위(최저점자 1위 · 최고점자 n위) · n = 평가군 인원 · s(n) = 인원 n일 때 z값들의 표준편차(17명 이상은 1). "
                    "60점 미만은 60점, 100점 초과는 100점으로 하고 소수점 셋째자리에서 반올림한다. "
                    "엑셀에서는 =ROUND(MEDIAN(60,80+7*NORM.S.INV((r-0.375)/(n+0.25))/s,100),3) 이다.", font=FN, align=WRAP)
us.merge_cells(start_row=ur, start_column=1, end_row=ur, end_column=6)
us.row_dimensions[ur].height = 46
ur += 1
put(us, "A%d" % ur, "조합이 먼저 밝히는 논점", font=B_SEC); ur += 1
put(us, "A%d" % ur,
    "제3호 후단의 「17명 이상은 s(n) = 1로 본다」를 그대로 적용하면 17명 이상 평가군의 표준편차가 "
    "7.00에 미달한다(n=17 → 6.51 · n=30 → 6.70 · n=40 → 6.76). 1차 평가군(부서 단위 3~16명)은 영향이 없으나 "
    "2차 평가군(본부 단위)은 통상 17명 이상이므로 영향을 받는다. 이 계산서는 규정 문언대로 구현하고 "
    "그 결과를 ⑤ 검증 시트 1-2절에 수치로 드러내 두었다. 해소하려면 제3호 후단을 삭제하고 s(n) 표를 "
    "인원 구간 전체로 확대하면 되며, ④ 시트가 이미 n = 40까지 계산해 두었다.", font=FN, align=WRAP)
us.merge_cells(start_row=ur, start_column=1, end_row=ur, end_column=6)
us.row_dimensions[ur].height = 74
widths(us, {"A": 30, "B": 34, "C": 30, "D": 24, "E": 24, "F": 24})

wb.save(OUT)
print("saved:", OUT)
